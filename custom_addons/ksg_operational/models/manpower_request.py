import io
import base64
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from odoo import models, fields, api, _
from odoo.exceptions import UserError

class KsgOperationalManpowerRequest(models.Model):
    _name = 'ksg.operational.manpower.request'
    _description = 'Permintaan Tenaga Kerja / Detail SDM'
    _inherit = ['mail.thread', 'mail.activity.mixin']
    _order = 'tanggal_pengajuan desc, id desc'

    name = fields.Char(string='Nomor Dokumen SDM', required=True, copy=False, default=lambda self: _('Draft SDM'))
    project_id = fields.Many2one('ksg.sales.project', string='Proyek Terkait', required=True, tracking=True)
    currency_id = fields.Many2one('res.currency', string='Mata Uang', default=lambda self: self.env.company.currency_id)

    tanggal_pengajuan = fields.Date(string='Tanggal Pengajuan', default=fields.Date.context_today, tracking=True)
    requested_by = fields.Many2one('res.users', string='Diajukan Oleh', default=lambda self: self.env.user)
    approver_id = fields.Many2one('res.users', string='Disetujui Oleh', readonly=True)

    # Field kriteria & kebutuhan waktu (kompatibilitas HC & view)
    kriteria = fields.Text(string='Kriteria / Kualifikasi Tenaga Kerja')
    kebutuhan_waktu = fields.Char(string='Kebutuhan Waktu / Periode Penugasan')
    jumlah_pekerja = fields.Integer(string='Jumlah Pekerja', compute='_compute_totals', store=True)

    # Parameter Dasar Kalkulasi (Sesuai Excel)
    dasar_umk = fields.Monetary(string='Dasar UMK (BPJS)', default=2400000.0, currency_field='currency_id', help='Contoh: UMK Rembang Rp 2.400.000')
    bpjs_kes_rate = fields.Float(string='Tarif BPJS Kesehatan', default=0.04)
    bpjs_tk_rate = fields.Float(string='Tarif BPJS TK', default=0.0689)
    seragam_biaya = fields.Monetary(string='Biaya Seragam / Bln', default=50000.0, currency_field='currency_id')
    overhead_rate = fields.Float(string='Tarif Overhead', default=0.02)
    fee_rate = fields.Float(string='Tarif Fee', default=0.02)

    # Tabel 1: Plotingan SDM
    plot_line_ids = fields.One2many('ksg.operational.manpower.plot.line', 'request_id', string='Plotingan SDM')
    
    # Tabel 2 & 3: Rincian Biaya & Unit Cost SDM
    cost_line_ids = fields.One2many('ksg.operational.manpower.cost.line', 'request_id', string='Rincian Biaya SDM')

    total_tenaga_kerja = fields.Integer(string='Total Personel (TK)', compute='_compute_totals', store=True)
    total_biaya_sdm_bulan = fields.Monetary(string='Total Biaya SDM / Bulan', compute='_compute_totals', store=True, currency_field='currency_id')
    total_biaya_sdm_tahun = fields.Monetary(string='Total Biaya SDM / Tahun', compute='_compute_totals', store=True, currency_field='currency_id')

    state = fields.Selection([
        ('draft', 'Draft'),
        ('submitted', 'Diajukan ke HC/Manajemen'),
        ('approved', 'Disetujui'),
        ('rejected', 'Ditolak')
    ], string='Status', default='draft', tracking=True)

    catatan = fields.Text(string='Catatan / Keterangan')

    @api.depends('plot_line_ids.total_tk', 'cost_line_ids.total_biaya')
    def _compute_totals(self):
        for rec in self:
            tot_tk = sum(rec.plot_line_ids.mapped('total_tk'))
            rec.total_tenaga_kerja = tot_tk
            rec.jumlah_pekerja = tot_tk
            rec.total_biaya_sdm_bulan = sum(rec.cost_line_ids.mapped('total_biaya'))
            rec.total_biaya_sdm_tahun = rec.total_biaya_sdm_bulan * 12

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            if vals.get('name', _('Draft SDM')) in [_('Draft SDM'), 'Draft SDM', False]:
                vals['name'] = self.env['ir.sequence'].next_by_code('ksg.operational.manpower') or _('Draft SDM')
        return super().create(vals_list)

    def action_generate_cost_lines(self):
        self.ensure_one()
        self.cost_line_ids.unlink()
        new_lines = []
        for p in self.plot_line_ids:
            if p.total_tk <= 0:
                continue
            is_tl = 'LEADER' in (p.jabatan or '').upper() or 'TL' in (p.jabatan or '').upper()
            gapok = 1900000.0 if is_tl else 1500000.0
            tunj = 300000.0 if is_tl else 0.0
            new_lines.append((0, 0, {
                'jabatan': p.jabatan,
                'unit_tk': p.total_tk,
                'gapok': gapok,
                'tunjangan': tunj,
            }))
        self.cost_line_ids = new_lines

    def action_submit(self):
        self.write({'state': 'submitted'})

    def action_approve(self):
        self.write({'state': 'approved', 'approver_id': self.env.user.id})

    def action_reject(self):
        self.write({'state': 'rejected'})

    def action_reset_draft(self):
        self.write({'state': 'draft'})

    def action_export_excel(self):
        self.ensure_one()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DETAIL SDM"

        font_title = Font(name="Calibri", size=11, bold=True)
        font_bold = Font(name="Calibri", size=10, bold=True)
        fill_header = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        border_thin = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

        ws['C2'] = "KOMPONEN BIAYA OPERASIONAL LANGSUNG"
        ws['C2'].font = font_title
        ws['C4'] = "PLOTINGAN SDM"
        ws['C4'].font = font_title

        headers_p = ["JABATAN", "SHIFT 1", "SHIFT 2", "LIBUR", "TOTAL"]
        for idx, h in enumerate(headers_p, start=3):
            c = ws.cell(row=5, column=idx, value=h)
            c.font = font_bold
            c.fill = fill_header
            c.border = border_thin
            c.alignment = Alignment(horizontal="center")

        ws.cell(row=6, column=4, value="06.00 - 14.00").alignment = Alignment(horizontal="center")
        ws.cell(row=6, column=5, value="14.00 - 22.00").alignment = Alignment(horizontal="center")

        r = 7
        for line in self.plot_line_ids:
            ws.cell(row=r, column=3, value=line.jabatan).border = border_thin
            ws.cell(row=r, column=4, value=line.shift_1).border = border_thin
            ws.cell(row=r, column=5, value=line.shift_2).border = border_thin
            ws.cell(row=r, column=6, value=line.libur).border = border_thin
            ws.cell(row=r, column=7, value=line.total_tk).border = border_thin
            r += 1

        ws.cell(row=r, column=3, value="TOTAL").font = font_bold
        ws.cell(row=r, column=4, value=sum(self.plot_line_ids.mapped('shift_1'))).font = font_bold
        ws.cell(row=r, column=5, value=sum(self.plot_line_ids.mapped('shift_2'))).font = font_bold
        ws.cell(row=r, column=6, value=sum(self.plot_line_ids.mapped('libur'))).font = font_bold
        ws.cell(row=r, column=7, value=self.total_tenaga_kerja).font = font_bold

        r += 3
        ws.cell(row=r, column=3, value="BIAYA SDM").font = font_title
        r += 1
        for idx, h in enumerate(["JABATAN", "UNIT", "UNIT COST", "TOTAL"], start=3):
            c = ws.cell(row=r, column=idx, value=h)
            c.font = font_bold
            c.fill = fill_header
            c.border = border_thin
            c.alignment = Alignment(horizontal="center")
        r += 1

        for c_line in self.cost_line_ids:
            ws.cell(row=r, column=3, value=c_line.jabatan).border = border_thin
            ws.cell(row=r, column=4, value=c_line.unit_tk).border = border_thin
            ws.cell(row=r, column=5, value=c_line.unit_cost).border = border_thin
            ws.cell(row=r, column=6, value=c_line.total_biaya).border = border_thin
            r += 1

        ws.cell(row=r, column=3, value="TOTAL BIAYA").font = font_bold
        ws.cell(row=r, column=4, value=self.total_tenaga_kerja).font = font_bold
        ws.cell(row=r, column=6, value=self.total_biaya_sdm_bulan).font = font_bold

        fp = io.BytesIO()
        wb.save(fp)
        data_base64 = base64.b64encode(fp.getvalue())

        attachment = self.env['ir.attachment'].create({
            'name': f'Permintaan_SDM_{self.name.replace("/", "_")}.xlsx',
            'type': 'binary',
            'datas': data_base64,
            'res_model': self._name,
            'res_id': self.id,
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        })
        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }


class KsgOperationalManpowerPlotLine(models.Model):
    _name = 'ksg.operational.manpower.plot.line'
    _description = 'Plotingan Shift SDM'

    request_id = fields.Many2one('ksg.operational.manpower.request', string='Dokumen SDM', ondelete='cascade')
    jabatan = fields.Char(string='Jabatan / Posisi', required=True)
    shift_1 = fields.Integer(string='Shift 1 (06.00-14.00)', default=0)
    shift_2 = fields.Integer(string='Shift 2 (14.00-22.00)', default=0)
    libur = fields.Integer(string='Libur', default=0)
    total_tk = fields.Integer(string='Total TK', compute='_compute_total_tk', store=True)

    @api.depends('shift_1', 'shift_2', 'libur')
    def _compute_total_tk(self):
        for line in self:
            line.total_tk = line.shift_1 + line.shift_2 + line.libur


class KsgOperationalManpowerCostLine(models.Model):
    _name = 'ksg.operational.manpower.cost.line'
    _description = 'Rincian Unit Cost SDM'

    request_id = fields.Many2one('ksg.operational.manpower.request', string='Dokumen SDM', ondelete='cascade')
    currency_id = fields.Many2one(related='request_id.currency_id')

    jabatan = fields.Char(string='Jabatan / Posisi', required=True)
    unit_tk = fields.Integer(string='Jumlah TK (Unit)', default=1)
    
    gapok = fields.Monetary(string='Gaji Pokok (GAPOK)', required=True, default=1500000.0, currency_field='currency_id')
    tunjangan = fields.Monetary(string='Tunjangan Jabatan', default=0.0, currency_field='currency_id')
    total_upah_tetap = fields.Monetary(string='Total Upah Tetap', compute='_compute_costs', store=True, currency_field='currency_id')

    bpjs_kes = fields.Monetary(string='BPJS Kesehatan', compute='_compute_costs', store=True, currency_field='currency_id')
    bpjs_tk = fields.Monetary(string='BPJS Ketenagakerjaan', compute='_compute_costs', store=True, currency_field='currency_id')
    seragam = fields.Monetary(string='Seragam', compute='_compute_costs', store=True, currency_field='currency_id')
    thr = fields.Monetary(string='THR / Bln', compute='_compute_costs', store=True, currency_field='currency_id')
    overhead = fields.Monetary(string='Overhead (2%)', compute='_compute_costs', store=True, currency_field='currency_id')
    fee = fields.Monetary(string='FEE (2%)', compute='_compute_costs', store=True, currency_field='currency_id')
    total_overhead = fields.Monetary(string='Total Overhead', compute='_compute_costs', store=True, currency_field='currency_id')

    unit_cost = fields.Monetary(string='Unit Cost SDM', compute='_compute_costs', store=True, currency_field='currency_id')
    total_biaya = fields.Monetary(string='Total Tagihan/Bulan', compute='_compute_costs', store=True, currency_field='currency_id')

    @api.depends('unit_tk', 'gapok', 'tunjangan', 'request_id.dasar_umk', 'request_id.bpjs_kes_rate', 
                 'request_id.bpjs_tk_rate', 'request_id.seragam_biaya', 'request_id.overhead_rate', 'request_id.fee_rate')
    def _compute_costs(self):
        for line in self:
            req = line.request_id
            umk = req.dasar_umk or 2400000.0
            line.total_upah_tetap = line.gapok + line.tunjangan

            line.bpjs_kes = umk * (req.bpjs_kes_rate or 0.04)
            line.bpjs_tk = umk * (req.bpjs_tk_rate or 0.0689)
            line.seragam = req.seragam_biaya or 50000.0
            line.thr = line.total_upah_tetap / 12.0
            line.overhead = line.total_upah_tetap * (req.overhead_rate or 0.02)
            line.fee = line.total_upah_tetap * (req.fee_rate or 0.02)

            line.total_overhead = line.bpjs_kes + line.bpjs_tk + line.seragam + line.thr + line.overhead + line.fee
            line.unit_cost = line.total_upah_tetap + line.total_overhead
            line.total_biaya = line.unit_cost * line.unit_tk