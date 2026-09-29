import io
import base64
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from odoo import models, fields, api, _

class KsgOperationalHpp(models.Model):
    _name = 'ksg.operational.hpp'
    _description = 'Kalkulasi Lembar HPP Operasional Proyek'
    _inherit = ['mail.thread', 'mail.activity.mixin']
    _order = 'tanggal desc, id desc'

    name = fields.Char(string='Nomor Dokumen HPP', required=True, copy=False, default=lambda self: _('Draft HPP'))
    project_id = fields.Many2one('ksg.sales.project', string='Proyek Terkait', required=True, tracking=True)
    currency_id = fields.Many2one('res.currency', string='Mata Uang', default=lambda self: self.env.company.currency_id)
    tanggal = fields.Date(string='Tanggal', default=fields.Date.context_today)

    manpower_id = fields.Many2one('ksg.operational.manpower.request', string='Dokumen SDM Terkait')
    supply_id = fields.Many2one('ksg.operational.supply.request', string='Dokumen Perlengkapan Terkait')
    source_status_info = fields.Char(string='Status Dokumen Terkait', compute='_compute_source_status_info')

    mk_bulan = fields.Integer(string='Masa Kontrak / MK (Bulan)', default=12)
    ovh_rate_hpp = fields.Float(string='Tarif OVH HPP', default=0.005, help='0.005 = 0.5% dari Uang Pokok')
    seragam_hpp = fields.Monetary(string='Seragam HPP / Org', default=40000.0, currency_field='currency_id')
    sistem_hpp = fields.Monetary(string='Sistem HPP / Org', default=10000.0, currency_field='currency_id')

    line_ids = fields.One2many('ksg.operational.hpp.line', 'hpp_id', string='Rincian Pos HPP')

    total_hpp_bulan = fields.Monetary(string='Total HPP / Bulan', compute='_compute_summary', store=True, currency_field='currency_id')
    total_hpp_tahun = fields.Monetary(string='Total HPP / Tahun', compute='_compute_summary', store=True, currency_field='currency_id')

    total_tagihan_bulan = fields.Monetary(string='Total Tagihan / Bulan', compute='_compute_summary', store=True, currency_field='currency_id')
    total_tagihan_tahun = fields.Monetary(string='Total Tagihan / Tahun', compute='_compute_summary', store=True, currency_field='currency_id')

    profit_bulan = fields.Monetary(string='Selisih (Profit) / Bulan', compute='_compute_summary', store=True, currency_field='currency_id')
    profit_tahun = fields.Monetary(string='Selisih (Profit) / Tahun', compute='_compute_summary', store=True, currency_field='currency_id')
    margin_percentage = fields.Float(string='Margin Keuntungan (%)', compute='_compute_summary', store=True, digits=(16, 4))

    alasan_penolakan = fields.Text(string='Alasan Penolakan', tracking=True)

    state = fields.Selection([
        ('draft', 'Draft'),
        ('approved', 'Disetujui Direktur'),
        ('rejected', 'Ditolak')
    ], string='Status', default='draft', tracking=True)

    @api.depends('project_id', 'manpower_id', 'supply_id')
    def _compute_source_status_info(self):
        for rec in self:
            sdm_txt = rec.manpower_id.name if rec.manpower_id else "Belum Dibuat"
            sup_txt = rec.supply_id.name if rec.supply_id else "Belum Dibuat"
            rec.source_status_info = f"SDM: {sdm_txt} | Perlengkapan: {sup_txt}"

    @api.depends('line_ids.jumlah_bulan', 'line_ids.hpp_tahun', 'line_ids.tagihan_bulan', 'line_ids.tagihan_tahun')
    def _compute_summary(self):
        for rec in self:
            rec.total_hpp_bulan = sum(rec.line_ids.mapped('jumlah_bulan'))
            rec.total_hpp_tahun = sum(rec.line_ids.mapped('hpp_tahun'))
            rec.total_tagihan_bulan = sum(rec.line_ids.mapped('tagihan_bulan'))
            rec.total_tagihan_tahun = sum(rec.line_ids.mapped('tagihan_tahun'))

            rec.profit_bulan = rec.total_tagihan_bulan - rec.total_hpp_bulan
            rec.profit_tahun = rec.total_tagihan_tahun - rec.total_hpp_tahun
            rec.margin_percentage = (rec.profit_tahun / rec.total_tagihan_tahun) if rec.total_tagihan_tahun > 0 else 0.0

    @api.onchange('project_id')
    def _onchange_project_id(self):
        if self.project_id:
            manpower = self.env['ksg.operational.manpower.request'].search([
                ('project_id', '=', self.project_id.id)
            ], order='id desc', limit=1)
            self.manpower_id = manpower.id if manpower else False

            supply = self.env['ksg.operational.supply.request'].search([
                ('project_id', '=', self.project_id.id)
            ], order='id desc', limit=1)
            self.supply_id = supply.id if supply else False

            self._sync_lines_data()
        else:
            self.manpower_id = False
            self.supply_id = False
            self.line_ids = [(5, 0, 0)]

    def _sync_lines_data(self):
        self.ensure_one()
        if self.project_id:
            if not self.manpower_id or self.manpower_id.project_id != self.project_id:
                manpower = self.env['ksg.operational.manpower.request'].search([
                    ('project_id', '=', self.project_id.id)
                ], order='id desc', limit=1)
                self.manpower_id = manpower.id if manpower else False

            if not self.supply_id or self.supply_id.project_id != self.project_id:
                supply = self.env['ksg.operational.supply.request'].search([
                    ('project_id', '=', self.project_id.id)
                ], order='id desc', limit=1)
                self.supply_id = supply.id if supply else False

        self.line_ids = [(5, 0, 0)]
        new_lines = []
        mk = self.mk_bulan or 12

        # 1. Baris SDM (Ambil Gapok asli dari Dokumen SDM)
        if self.manpower_id:
            for cost in self.manpower_id.cost_line_ids:
                uang_pokok = cost.gapok
                tunj = cost.tunjangan
                new_lines.append((0, 0, {
                    'pos_type': 'sdm',
                    'bagian': cost.jabatan,
                    'mk': mk,
                    'tk': cost.unit_tk,
                    'uang_pokok': uang_pokok,
                    'tunj': tunj,
                    'seragam': self.seragam_hpp,
                    'sistem': self.sistem_hpp,
                    'bpjstk': cost.bpjs_tk,
                    'bpjkes': cost.bpjs_kes,
                    'tagihan_bulan': cost.total_biaya,
                }))

        # 2. Baris ALAT DAN BAHAN
        if self.supply_id:
            alat_bahan_bln = self.supply_id.total_alat_bahan_bulan
            new_lines.append((0, 0, {
                'pos_type': 'alat_bahan',
                'bagian': 'ALAT DAN BAHAN',
                'mk': mk,
                'tk': 0,
                'jumlah_bulan': alat_bahan_bln,
                'hpp_tahun': alat_bahan_bln * mk,
                'tagihan_bulan': alat_bahan_bln,
                'tagihan_tahun': alat_bahan_bln * mk,
            }))

            # 3. Baris JASA
            jasa_bln = self.supply_id.total_jasa_bulan
            new_lines.append((0, 0, {
                'pos_type': 'jasa',
                'bagian': 'JASA',
                'mk': mk,
                'tk': 0,
                'jumlah_bulan': jasa_bln,
                'hpp_tahun': jasa_bln * mk,
                'tagihan_bulan': jasa_bln,
                'tagihan_tahun': jasa_bln * mk,
            }))

        self.line_ids = new_lines

    def action_sync_from_sdm_and_supply(self):
        self.ensure_one()
        self._sync_lines_data()

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            if vals.get('name', _('Draft HPP')) in [_('Draft HPP'), 'Draft HPP', False]:
                vals['name'] = self.env['ir.sequence'].next_by_code('ksg.operational.hpp') or _('Draft HPP')
            p_id = vals.get('project_id')
            if p_id:
                if not vals.get('manpower_id'):
                    mp = self.env['ksg.operational.manpower.request'].search([
                        ('project_id', '=', p_id)
                    ], order='id desc', limit=1)
                    if mp:
                        vals['manpower_id'] = mp.id
                if not vals.get('supply_id'):
                    sp = self.env['ksg.operational.supply.request'].search([
                        ('project_id', '=', p_id)
                    ], order='id desc', limit=1)
                    if sp:
                        vals['supply_id'] = sp.id
        recs = super().create(vals_list)
        for rec in recs:
            if not rec.line_ids and (rec.manpower_id or rec.supply_id):
                rec._sync_lines_data()
        return recs

    def action_approve(self):
        self.write({'state': 'approved'})

    def action_reset_draft(self):
        self.write({'state': 'draft'})

    def action_reject_wizard(self):
        self.ensure_one()
        return {
            'name': _('Tolak Lembar HPP'),
            'type': 'ir.actions.act_window',
            'res_model': 'ksg.operational.reject.wizard',
            'view_mode': 'form',
            'target': 'new',
            'context': {
                'default_res_model': self._name,
                'default_res_id': self.id,
            }
        }

    def action_export_excel(self):
        self.ensure_one()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "HPP"

        font_bold = Font(name="Calibri", size=10, bold=True)
        font_title = Font(name="Calibri", size=11, bold=True)
        fill_header = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        fill_yellow = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        fill_green = PatternFill(start_color="00B050", end_color="00B050", fill_type="solid")
        border_thin = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

        proj_name = self.project_id.display_name if self.project_id else ''
        ws['B1'] = f"LEMBAR HPP OPERASIONAL: {proj_name}"
        ws['B1'].font = font_title

        headers = ["NO", "BAGIAN", "MK", "TK", "UANG POKOK", "TUNJ", "THR", "SERAGAM", "OVH", "SISTEM", "BPJSTK", "BPJKES", "EX PENGELUARAN", "JUMLAH/BULAN", "HPP/Tahun", "TAGIHAN/Bulan", "TAGIHAN/TAHUN"]

        for idx, h in enumerate(headers, start=2):
            c = ws.cell(row=3, column=idx, value=h)
            c.font = font_bold
            c.fill = fill_header
            c.border = border_thin
            c.alignment = Alignment(horizontal="center")

        r = 5
        no = 1
        for l in self.line_ids:
            ws.cell(row=r, column=2, value=no).border = border_thin
            ws.cell(row=r, column=3, value=l.bagian).border = border_thin
            ws.cell(row=r, column=4, value=l.mk).border = border_thin
            ws.cell(row=r, column=5, value=l.tk or '').border = border_thin
            ws.cell(row=r, column=6, value=l.uang_pokok or '').border = border_thin
            ws.cell(row=r, column=7, value=l.tunj or '').border = border_thin
            ws.cell(row=r, column=8, value=l.thr or '').border = border_thin
            ws.cell(row=r, column=9, value=l.seragam or '').border = border_thin
            ws.cell(row=r, column=10, value=l.ovh or '').border = border_thin
            ws.cell(row=r, column=11, value=l.sistem or '').border = border_thin
            ws.cell(row=r, column=12, value=l.bpjstk or '').border = border_thin
            ws.cell(row=r, column=13, value=l.bpjkes or '').border = border_thin
            ws.cell(row=r, column=14, value=l.ex_pengeluaran or '').border = border_thin
            ws.cell(row=r, column=15, value=l.jumlah_bulan).border = border_thin
            ws.cell(row=r, column=16, value=l.hpp_tahun).border = border_thin
            ws.cell(row=r, column=17, value=l.tagihan_bulan).border = border_thin
            ws.cell(row=r, column=18, value=l.tagihan_tahun).border = border_thin
            no += 1
            r += 1

        c_tot = ws.cell(row=r, column=3, value="TOTAL")
        c_tot.font = font_bold
        for col_idx, val in [(15, self.total_hpp_bulan), (16, self.total_hpp_tahun), (17, self.total_tagihan_bulan), (18, self.total_tagihan_tahun)]:
            c = ws.cell(row=r, column=col_idx, value=val)
            c.font = font_bold
            c.fill = fill_green
            c.border = border_thin
        r += 1

        ws.cell(row=r, column=3, value="SELISIH (PROFIT)").font = font_bold
        c_prof_m = ws.cell(row=r, column=17, value=self.profit_bulan)
        c_prof_m.font = font_bold
        c_prof_m.fill = fill_yellow
        c_prof_m.border = border_thin

        c_prof_y = ws.cell(row=r, column=18, value=self.profit_tahun)
        c_prof_y.font = font_bold
        c_prof_y.fill = fill_yellow
        c_prof_y.border = border_thin
        r += 1

        ws.cell(row=r, column=3, value="MARGIN %").font = font_bold
        pct_str = f"{(self.margin_percentage * 100):.4f}%"
        c_pct_m = ws.cell(row=r, column=17, value=pct_str)
        c_pct_m.font = font_bold
        c_pct_m.fill = fill_yellow
        c_pct_m.border = border_thin

        c_pct_y = ws.cell(row=r, column=18, value=pct_str)
        c_pct_y.font = font_bold
        c_pct_y.fill = fill_yellow
        c_pct_y.border = border_thin

        fp = io.BytesIO()
        wb.save(fp)
        data_base64 = base64.b64encode(fp.getvalue())

        attachment = self.env['ir.attachment'].create({
            'name': f'Lembar_HPP_{self.name.replace("/", "_")}.xlsx',
            'type': 'binary',
            'datas': data_base64,
            'res_model': self._name,
            'res_id': self.id,
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        })
        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }


class KsgOperationalHppLine(models.Model):
    _name = 'ksg.operational.hpp.line'
    _description = 'Baris Pos HPP'

    hpp_id = fields.Many2one('ksg.operational.hpp', string='Dokumen HPP', ondelete='cascade')
    currency_id = fields.Many2one(related='hpp_id.currency_id')

    pos_type = fields.Selection([
        ('sdm', 'Tenaga Kerja (SDM)'),
        ('alat_bahan', 'Alat dan Bahan'),
        ('jasa', 'Jasa')
    ], string='Jenis Pos', required=True, default='sdm')

    bagian = fields.Char(string='Bagian / Pos Pekerjaan', required=True)
    mk = fields.Integer(string='MK (Bln)', default=12)
    tk = fields.Integer(string='TK (Org)', default=1)

    uang_pokok = fields.Monetary(string='Uang Pokok', currency_field='currency_id')
    tunj = fields.Monetary(string='Tunjangan', currency_field='currency_id')
    thr = fields.Monetary(string='THR', compute='_compute_sdm_costs', store=True, currency_field='currency_id')
    seragam = fields.Monetary(string='Seragam', currency_field='currency_id')
    ovh = fields.Monetary(string='OVH', compute='_compute_sdm_costs', store=True, currency_field='currency_id')
    sistem = fields.Monetary(string='Sistem', currency_field='currency_id')
    bpjstk = fields.Monetary(string='BPJS TK', currency_field='currency_id')
    bpjkes = fields.Monetary(string='BPJS Kes', currency_field='currency_id')

    ex_pengeluaran = fields.Monetary(string='Ex Pengeluaran', compute='_compute_sdm_costs', store=True, currency_field='currency_id')
    jumlah_bulan = fields.Monetary(string='Jumlah / Bulan (HPP)', compute='_compute_sdm_costs', store=True, readonly=False, currency_field='currency_id')
    hpp_tahun = fields.Monetary(string='HPP / Tahun', compute='_compute_line_totals', store=True, currency_field='currency_id')

    tagihan_bulan = fields.Monetary(string='Tagihan / Bulan', currency_field='currency_id')
    tagihan_tahun = fields.Monetary(string='Tagihan / Tahun', compute='_compute_line_totals', store=True, currency_field='currency_id')

    @api.depends('pos_type', 'mk', 'tk', 'uang_pokok', 'tunj', 'seragam', 'sistem', 'bpjstk', 'bpjkes', 'hpp_id.ovh_rate_hpp')
    def _compute_sdm_costs(self):
        for line in self:
            if line.pos_type == 'sdm':
                rate_ovh = line.hpp_id.ovh_rate_hpp or 0.005
                line.thr = (line.uang_pokok + line.tunj) / 12.0
                line.ovh = line.uang_pokok * rate_ovh
                line.ex_pengeluaran = line.uang_pokok + line.tunj + line.thr + line.seragam + line.ovh + line.sistem + line.bpjstk + line.bpjkes
                line.jumlah_bulan = line.ex_pengeluaran * line.tk
            else:
                line.thr = 0.0
                line.ovh = 0.0
                line.ex_pengeluaran = 0.0

    @api.depends('jumlah_bulan', 'tagihan_bulan', 'mk')
    def _compute_line_totals(self):
        for line in self:
            m = line.mk or 12
            line.hpp_tahun = line.jumlah_bulan * m
            line.tagihan_tahun = line.tagihan_bulan * m