import io
import base64
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from odoo import models, fields, api, _
from odoo.exceptions import UserError

class KsgOperationalSupplyRequest(models.Model):
    _name = 'ksg.operational.supply.request'
    _description = 'Permintaan Perlengkapan, Chemical & Non Listed'
    _inherit = ['mail.thread', 'mail.activity.mixin']
    _order = 'tanggal_pengajuan desc, id desc'

    name = fields.Char(string='Nomor Dokumen', required=True, copy=False, default=lambda self: _('Draft Perlengkapan'))
    project_id = fields.Many2one('ksg.sales.project', string='Proyek Terkait', required=True, tracking=True)
    currency_id = fields.Many2one('res.currency', string='Mata Uang', default=lambda self: self.env.company.currency_id)

    tanggal_pengajuan = fields.Date(string='Tanggal Pengajuan', default=fields.Date.context_today, tracking=True)
    durasi_bulan = fields.Integer(string='Durasi Proyek (Bulan)', default=12)
    requested_by = fields.Many2one('res.users', string='Diajukan Oleh', default=lambda self: self.env.user)
    approver_id = fields.Many2one('res.users', string='Disetujui Oleh', readonly=True)

    line_ids = fields.One2many('ksg.operational.supply.line', 'request_id', string='Daftar Kebutuhan')

    # Rekapitulasi Otomatis Sub A, B, C
    total_perlengkapan_tahun = fields.Monetary(string='Sub A: Perlengkapan / Tahun', compute='_compute_totals', store=True, currency_field='currency_id')
    total_perlengkapan_bulan = fields.Monetary(string='Sub A: Perlengkapan / Bulan', compute='_compute_totals', store=True, currency_field='currency_id')

    total_chemical_tahun = fields.Monetary(string='Sub B: Chemical / Tahun', compute='_compute_totals', store=True, currency_field='currency_id')
    total_chemical_bulan = fields.Monetary(string='Sub B: Chemical / Bulan', compute='_compute_totals', store=True, currency_field='currency_id')

    total_non_listed_tahun = fields.Monetary(string='Sub C: Non Listed / Tahun', compute='_compute_totals', store=True, currency_field='currency_id')
    total_non_listed_bulan = fields.Monetary(string='Sub C: Non Listed / Bulan', compute='_compute_totals', store=True, currency_field='currency_id')

    # Pengelompokan untuk Lembar HPP
    total_alat_bahan_bulan = fields.Monetary(string='Alat & Bahan (A+B) / Bulan', compute='_compute_totals', store=True, currency_field='currency_id')
    total_alat_bahan_tahun = fields.Monetary(string='Alat & Bahan (A+B) / Tahun', compute='_compute_totals', store=True, currency_field='currency_id')
    total_jasa_bulan = fields.Monetary(string='Jasa (C) / Bulan', compute='_compute_totals', store=True, currency_field='currency_id')
    total_jasa_tahun = fields.Monetary(string='Jasa (C) / Tahun', compute='_compute_totals', store=True, currency_field='currency_id')

    grand_total_bulan = fields.Monetary(string='Total Tagihan / Bulan', compute='_compute_totals', store=True, currency_field='currency_id')
    grand_total_tahun = fields.Monetary(string='Total Tagihan / Tahun', compute='_compute_totals', store=True, currency_field='currency_id')

    state = fields.Selection([
        ('draft', 'Draft'),
        ('submitted', 'Diajukan ke Direktur'),
        ('approved', 'Disetujui'),
        ('rejected', 'Ditolak')
    ], string='Status', default='draft', tracking=True)

    catatan = fields.Text(string='Catatan / Keterangan')

    @api.depends('line_ids.total_harga', 'durasi_bulan')
    def _compute_totals(self):
        for rec in self:
            m = rec.durasi_bulan if rec.durasi_bulan > 0 else 12
            lines_a = rec.line_ids.filtered(lambda l: l.sub_jenis == 'A')
            lines_b = rec.line_ids.filtered(lambda l: l.sub_jenis == 'B')
            lines_c = rec.line_ids.filtered(lambda l: l.sub_jenis == 'C')

            rec.total_perlengkapan_tahun = sum(lines_a.mapped('total_harga'))
            rec.total_perlengkapan_bulan = rec.total_perlengkapan_tahun / m

            rec.total_chemical_tahun = sum(lines_b.mapped('total_harga'))
            rec.total_chemical_bulan = rec.total_chemical_tahun / m

            rec.total_non_listed_tahun = sum(lines_c.mapped('total_harga'))
            rec.total_non_listed_bulan = rec.total_non_listed_tahun / m

            rec.total_alat_bahan_bulan = rec.total_perlengkapan_bulan + rec.total_chemical_bulan
            rec.total_alat_bahan_tahun = rec.total_perlengkapan_tahun + rec.total_chemical_tahun
            rec.total_jasa_bulan = rec.total_non_listed_bulan
            rec.total_jasa_tahun = rec.total_non_listed_tahun

            rec.grand_total_bulan = rec.total_alat_bahan_bulan + rec.total_jasa_bulan
            rec.grand_total_tahun = rec.total_alat_bahan_tahun + rec.total_jasa_tahun

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            if vals.get('name', _('Draft Perlengkapan')) in [_('Draft Perlengkapan'), 'Draft Perlengkapan', False]:
                vals['name'] = self.env['ir.sequence'].next_by_code('ksg.operational.supply') or _('Draft Perlengkapan')
        return super().create(vals_list)

    def action_submit(self):
        self.write({'state': 'submitted'})

    def action_approve(self):
        self.write({'state': 'approved', 'approver_id': self.env.user.id})

    def action_reject(self):
        self.write({'state': 'rejected'})

    def action_reset_draft(self):
        self.write({'state': 'draft'})

    def action_export_excel(self):
        self.ensure_one()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Perlengkapan n Chemical"

        font_sec = Font(name="Calibri", size=11, bold=True)
        font_bold = Font(name="Calibri", size=10, bold=True)
        fill_header = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        border_thin = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

        headers = ["No", "Item", "Qty", "UoM Periodik", "Harga Satuan HPP", "Total Harga"]

        current_row = 2
        for code, label in [('A', 'PERLENGKAPAN'), ('B', 'CHEMICAL'), ('C', 'NON LISTED')]:
            ws.cell(row=current_row, column=2, value=code).font = font_sec
            ws.cell(row=current_row, column=3, value=label).font = font_sec
            current_row += 2

            for idx, h in enumerate(headers, start=2):
                c = ws.cell(row=current_row, column=idx, value=h)
                c.font = font_bold
                c.fill = fill_header
                c.border = border_thin
                c.alignment = Alignment(horizontal="center")
            current_row += 1

            sub_lines = self.line_ids.filtered(lambda l: l.sub_jenis == code)
            sub_total = 0.0
            idx_num = 1
            for l in sub_lines:
                ws.cell(row=current_row, column=2, value=idx_num).border = border_thin
                ws.cell(row=current_row, column=3, value=l.item).border = border_thin
                ws.cell(row=current_row, column=4, value=l.qty).border = border_thin
                ws.cell(row=current_row, column=5, value=l.uom_periodik or 'Ea/Tahun').border = border_thin
                ws.cell(row=current_row, column=6, value=l.harga_satuan).border = border_thin
                ws.cell(row=current_row, column=7, value=l.total_harga).border = border_thin
                sub_total += l.total_harga
                idx_num += 1
                current_row += 1

            ws.cell(row=current_row, column=2, value=f"Total Biaya Kebutuhan {label} dalam 1 Tahun").font = font_bold
            ws.cell(row=current_row, column=7, value=sub_total).font = font_bold
            current_row += 2

        # Rekapitulasi Tabel Akhir
        ws.cell(row=current_row, column=2, value="NO").font = font_bold
        ws.cell(row=current_row, column=3, value="Item").font = font_bold
        ws.cell(row=current_row, column=4, value="Qty").font = font_bold
        ws.cell(row=current_row, column=5, value="SATUAN").font = font_bold
        ws.cell(row=current_row, column=6, value="BIAYA PERBULAN").font = font_bold
        ws.cell(row=current_row, column=7, value="DURASI").font = font_bold
        ws.cell(row=current_row, column=8, value="BIAYA PERTAHUN").font = font_bold
        current_row += 1

        rekap_data = [
            ('A', 'PERLENGKAPAN', self.total_perlengkapan_bulan, self.durasi_bulan, self.total_perlengkapan_tahun),
            ('B', 'CHEMICAL', self.total_chemical_bulan, self.durasi_bulan, self.total_chemical_tahun),
            ('C', 'NON LISTED', self.total_non_listed_bulan, self.durasi_bulan, self.total_non_listed_tahun),
        ]
        for r_code, r_item, r_bln, r_dur, r_thn in rekap_data:
            ws.cell(row=current_row, column=2, value=r_code).border = border_thin
            ws.cell(row=current_row, column=3, value=r_item).border = border_thin
            ws.cell(row=current_row, column=4, value=1).border = border_thin
            ws.cell(row=current_row, column=5, value="Lot").border = border_thin
            ws.cell(row=current_row, column=6, value=r_bln).border = border_thin
            ws.cell(row=current_row, column=7, value=r_dur).border = border_thin
            ws.cell(row=current_row, column=8, value=r_thn).border = border_thin
            current_row += 1

        ws.cell(row=current_row, column=3, value="TOTAL TAGIHAN 1 BULAN").font = font_bold
        ws.cell(row=current_row, column=6, value=self.grand_total_bulan).font = font_bold
        current_row += 1
        ws.cell(row=current_row, column=3, value="TOTAL TAGIHAN 1 TAHUN").font = font_bold
        ws.cell(row=current_row, column=8, value=self.grand_total_tahun).font = font_bold

        fp = io.BytesIO()
        wb.save(fp)
        data_base64 = base64.b64encode(fp.getvalue())

        attachment = self.env['ir.attachment'].create({
            'name': f'Permintaan_Perlengkapan_{self.name.replace("/", "_")}.xlsx',
            'type': 'binary',
            'datas': data_base64,
            'res_model': self._name,
            'res_id': self.id,
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        })
        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }


class KsgOperationalSupplyLine(models.Model):
    _name = 'ksg.operational.supply.line'
    _description = 'Item Permintaan Perlengkapan'

    request_id = fields.Many2one('ksg.operational.supply.request', string='Dokumen', ondelete='cascade')
    currency_id = fields.Many2one(related='request_id.currency_id')

    sub_jenis = fields.Selection([
        ('A', 'A. PERLENGKAPAN'),
        ('B', 'B. CHEMICAL'),
        ('C', 'C. NON LISTED')
    ], string='Sub-Jenis', required=True, default='A')

    item = fields.Char(string='Nama Barang / Uraian Pekerjaan', required=True)
    qty = fields.Float(string='Volume / Qty', required=True, default=1.0)
    uom_periodik = fields.Char(string='Satuan', default='Ea/Tahun')
    harga_satuan = fields.Monetary(string='Harga Satuan (Rp)', required=True, currency_field='currency_id')
    total_harga = fields.Monetary(string='Total Harga', compute='_compute_line_totals', store=True, currency_field='currency_id')
    biaya_per_bulan = fields.Monetary(string='Biaya / Bulan', compute='_compute_line_totals', store=True, currency_field='currency_id')
    keterangan = fields.Char(string='Keterangan')

    @api.depends('qty', 'harga_satuan', 'request_id.durasi_bulan')
    def _compute_line_totals(self):
        for line in self:
            m = line.request_id.durasi_bulan if (line.request_id and line.request_id.durasi_bulan > 0) else 12
            line.total_harga = line.qty * line.harga_satuan
            line.biaya_per_bulan = line.total_harga / m