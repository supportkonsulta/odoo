import io
import base64
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from odoo import models, fields, api, _
from odoo.exceptions import ValidationError, UserError

class KsgSalesWorker(models.Model):
    _name = 'ksg.sales.worker'
    _description = 'Master Data Tenaga Kerja'
    _inherit = ['mail.thread', 'mail.activity.mixin']

    name = fields.Char(string='Nama Lengkap', required=True, tracking=True)
    nik = fields.Char(string='NIK / No. KTP', tracking=True)
    gender = fields.Selection([
        ('male', 'Laki-laki'),
        ('female', 'Perempuan')
    ], string='Jenis Kelamin', default='male')
    pendidikan_terakhir = fields.Selection([
        ('sd', 'SD / Sederajat'),
        ('smp', 'SMP / Sederajat'),
        ('sma', 'SMA / SMK'),
        ('d3', 'Diploma (D3)'),
        ('s1', 'Sarjana (S1)'),
        ('s2', 'Magister (S2)')
    ], string='Pendidikan Terakhir', default='sma')
    posisi = fields.Char(string='Jabatan / Posisi Kerja', tracking=True)
    no_hp = fields.Char(string='No. Telepon / WhatsApp')
    alamat = fields.Text(string='Alamat Domisili')
    status_kontrak = fields.Selection([
        ('pkwt', 'PKWT (Kontrak Waktu Tertentu)'),
        ('pkwtt', 'PKWTT (Tetap)'),
        ('magang', 'Magang / Praktik'),
        ('freelance', 'Harian Lepas / Freelance')
    ], string='Status Kepegawaian', default='pkwt', tracking=True)
    status_resign = fields.Boolean(string='Status Resign', default=False, tracking=True)
    kompensasi_locked = fields.Boolean(string='Kompensasi Dikunci (Exit Clearance)', default=False, tracking=True)


class KsgOperationalManpowerRequestInherit(models.Model):
    _inherit = 'ksg.operational.manpower.request'

    cost_sheet_id = fields.Many2one(
        'ksg.hc.manpower.cost.sheet', 
        string='Dokumen Perhitungan SDM (HC)', 
        compute='_compute_cost_sheet_id'
    )

    def _compute_cost_sheet_id(self):
        for rec in self:
            sheet = self.env['ksg.hc.manpower.cost.sheet'].search([
                ('operational_request_id', '=', rec.id)
            ], limit=1)
            rec.cost_sheet_id = sheet.id if sheet else False

    def action_submit(self):
        res = super().action_submit()
        for rec in self:
            sheet = self.env['ksg.hc.manpower.cost.sheet'].search([
                ('operational_request_id', '=', rec.id)
            ], limit=1)
            if not sheet:
                sheet = self.env['ksg.hc.manpower.cost.sheet'].create({
                    'project_id': rec.project_id.id,
                    'operational_request_id': rec.id,
                    'tanggal': fields.Date.context_today(rec),
                })
            sheet.action_pull_from_operational()
        return res


class KsgHcManpowerCostSheet(models.Model):
    _name = 'ksg.hc.manpower.cost.sheet'
    _description = 'Base Dokumen Perhitungan SDM Proyek (Detail SDM untuk HPP Penjualan)'
    _inherit = ['mail.thread', 'mail.activity.mixin']
    _order = 'tanggal desc, id desc'

    name = fields.Char(string='Nomor Dokumen Perhitungan SDM', required=True, copy=False, default=lambda self: _('Draft SDM Cost'))
    project_id = fields.Many2one('ksg.sales.project', string='Proyek Terkait', required=True, tracking=True)
    operational_request_id = fields.Many2one('ksg.operational.manpower.request', string='Dokumen Permintaan Operasional', tracking=True)
    tanggal = fields.Date(string='Tanggal Perhitungan', default=fields.Date.context_today)
    currency_id = fields.Many2one('res.currency', string='Mata Uang', default=lambda self: self.env.company.currency_id)

    dasar_umk = fields.Monetary(string='Dasar UMK (BPJS)', default=2400000.0, currency_field='currency_id')
    keterangan_umk = fields.Char(string='Keterangan Acuan UMK', default='1. Nilai BPJS mengunakan UMK Rembang')
    bpjs_kes_rate = fields.Float(string='Tarif BPJS Kesehatan', default=0.04)
    bpjs_tk_rate = fields.Float(string='Tarif BPJS TK', default=0.0689)
    seragam_biaya = fields.Monetary(string='Biaya Seragam / Bln', default=50000.0, currency_field='currency_id')

    line_ids = fields.One2many('ksg.hc.manpower.cost.sheet.line', 'cost_sheet_id', string='Rincian Pos SDM')

    total_tenaga_kerja = fields.Integer(string='Total Personil (TK)', compute='_compute_totals', store=True)
    total_biaya_sdm_bulan = fields.Monetary(string='Total Biaya SDM / Bulan', compute='_compute_totals', store=True, currency_field='currency_id')
    total_biaya_sdm_tahun = fields.Monetary(string='Total Biaya SDM / Tahun', compute='_compute_totals', store=True, currency_field='currency_id')

    catatan_hc = fields.Text(string='Catatan Evaluasi HC')
    alasan_penolakan = fields.Text(string='Alasan Penolakan ke Operasional')

    state = fields.Selection([
        ('draft', 'Draft (HC Review & Kalkulasi)'),
        ('approved', 'Disetujui & Disahkan (Siap untuk HPP)'),
        ('rejected', 'Ditolak ke Operasional')
    ], string='Status Dokumen', default='draft', tracking=True)

    @api.depends('line_ids.unit_tk', 'line_ids.total_biaya')
    def _compute_totals(self):
        for rec in self:
            rec.total_tenaga_kerja = sum(rec.line_ids.mapped('unit_tk'))
            rec.total_biaya_sdm_bulan = sum(rec.line_ids.mapped('total_biaya'))
            rec.total_biaya_sdm_tahun = rec.total_biaya_sdm_bulan * 12

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            if vals.get('name', _('Draft SDM Cost')) in [_('Draft SDM Cost'), 'Draft SDM Cost', False]:
                vals['name'] = self.env['ir.sequence'].next_by_code('ksg.hc.manpower.cost.sheet') or _('Draft SDM Cost')
        return super().create(vals_list)

    def write(self, vals):
        res = super().write(vals)
        if 'state' in vals and vals['state'] == 'rejected':
            for rec in self:
                if rec.operational_request_id:
                    rec.operational_request_id.write({
                        'state': 'rejected',
                        'alasan_penolakan': rec.alasan_penolakan or vals.get('alasan_penolakan')
                    })
        return res

    def action_pull_from_operational(self):
        self.ensure_one()
        req = self.operational_request_id
        if not req and self.project_id:
            req = self.env['ksg.operational.manpower.request'].search([
                ('project_id', '=', self.project_id.id)
            ], order='id desc', limit=1)
            self.operational_request_id = req.id if req else False

        if not req:
            raise UserError(_("Tidak ditemukan dokumen Permintaan Tenaga Kerja Operasional untuk proyek ini."))

        self.line_ids.unlink()
        new_lines = []
        umk = self.dasar_umk or 2400000.0
        kes_rate = self.bpjs_kes_rate or 0.04
        tk_rate = self.bpjs_tk_rate or 0.0689
        seragam = self.seragam_biaya or 50000.0

        for plot in req.plot_line_ids:
            is_tl = 'LEADER' in (plot.jabatan or '').upper() or 'TL' in (plot.jabatan or '').upper()
            gapok = 1900000.0 if is_tl else 1500000.0
            tunjangan = 300000.0 if is_tl else 0.0
            new_lines.append((0, 0, {
                'jabatan': plot.jabatan,
                'shift_1': plot.shift_1,
                'shift_2': plot.shift_2,
                'libur': plot.libur,
                'unit_tk': plot.total_tk,
                'gapok': gapok,
                'tunjangan': tunjangan,
                'bpjs_kes': umk * kes_rate,
                'bpjs_tk': umk * tk_rate,
                'seragam': seragam,
            }))
        self.line_ids = new_lines

    def action_approve(self):
        self.ensure_one()
        if not self.line_ids:
            raise UserError(_("Rincian biaya SDM belum ada. Silakan klik 'Tarik / Refresh Data Operasional' terlebih dahulu."))
        self.write({'state': 'approved'})
        if self.operational_request_id:
            self.operational_request_id.write({'state': 'approved'})

    def action_reject_to_operational(self):
        self.ensure_one()
        return {
            'name': _('Tolak Permintaan SDM ke Operasional'),
            'type': 'ir.actions.act_window',
            'res_model': 'ksg.operational.reject.wizard',
            'view_mode': 'form',
            'target': 'new',
            'context': {
                'default_res_model': self._name,
                'default_res_id': self.id,
            }
        }

    def action_reset_draft(self):
        self.write({'state': 'draft'})
        if self.operational_request_id and self.operational_request_id.state == 'approved':
            self.operational_request_id.write({'state': 'submitted'})

    def action_export_excel(self):
        self.ensure_one()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DETAIL SDM"
        ws.views.sheetView[0].showGridLines = True

        # Lebar Kolom Persis Seperti Aslinya
        ws.column_dimensions['B'].width = 5
        ws.column_dimensions['C'].width = 43
        ws.column_dimensions['D'].width = 11
        ws.column_dimensions['E'].width = 14
        ws.column_dimensions['F'].width = 11
        ws.column_dimensions['G'].width = 13
        ws.column_dimensions['H'].width = 14
        ws.column_dimensions['I'].width = 5

        # Styles
        f_title = Font(name="Calibri", size=12, bold=True)
        f_sec_hdr = Font(name="Calibri", size=11, bold=True)
        f_sub_hdr = Font(name="Calibri", size=10, bold=True)
        f_reg_11 = Font(name="Calibri", size=11, bold=False)
        f_bold_11 = Font(name="Calibri", size=11, bold=True)
        f_reg_9 = Font(name="Calibri", size=9, bold=False)
        f_bold_9 = Font(name="Calibri", size=9, bold=True)

        fill_yellow = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")

        med_side = Side(style='medium')
        thin_side = Side(style='thin')
        double_side = Side(style='double')

        num_fmt_curr = '_-* #,##0_-;\\-* #,##0_-;_-* "-"_-;_-@_-'
        num_fmt_int = '_-* #,##0_-;\\-* #,##0_-;_-* "-"_-;_-@_-'

        # Judul Dokumen (Row 2)
        ws['C2'] = "KOMPONEN BIAYA OPERASIONAL LANGSUNG"
        ws['C2'].font = f_title

        lines = self.line_ids
        num_pos = max(len(lines), 1)

        # ----------------------------------------------------
        # 1. SECTION PLOTINGAN SDM
        # ----------------------------------------------------
        ws['C4'] = "PLOTINGAN SDM"
        ws['C4'].font = f_sec_hdr
        for col_idx in range(3, 9):
            c = ws.cell(4, col_idx)
            c.fill = fill_yellow
            c.border = Border(top=med_side, left=med_side if col_idx == 3 else None, right=med_side if col_idx == 8 else None)

        ws['D5'] = "SHIFT 1"
        ws['D5'].font = f_sub_hdr
        ws['E5'] = "SHIFT 2"
        ws['E5'].font = f_sub_hdr
        ws['E5'].alignment = Alignment(horizontal="center")
        ws['G5'] = "LIBUR"
        ws['G5'].font = f_sub_hdr
        ws['G5'].alignment = Alignment(horizontal="center")
        ws['H5'] = "TOTAL"
        ws['H5'].font = f_sub_hdr
        ws['H5'].alignment = Alignment(horizontal="center")
        ws.cell(5, 3).border = Border(left=med_side)
        ws.cell(5, 8).border = Border(right=med_side)

        ws['D6'] = "06.00 - 14.00"
        ws['D6'].font = f_reg_9
        ws['E6'] = "14.00 - 22.00"
        ws['E6'].font = f_reg_9
        ws['E6'].alignment = Alignment(horizontal="center")
        ws.cell(6, 3).border = Border(left=med_side)
        ws.cell(6, 8).border = Border(right=med_side)

        plot_start_row = 7
        for idx, l in enumerate(lines):
            curr_r = plot_start_row + idx
            ws.cell(curr_r, 3, l.jabatan).font = f_reg_11
            ws.cell(curr_r, 3).border = Border(left=med_side)

            c_d = ws.cell(curr_r, 4, l.shift_1)
            c_d.font = f_reg_9
            c_d.alignment = Alignment(horizontal="right")

            c_e = ws.cell(curr_r, 5, l.shift_2)
            c_e.font = f_reg_9
            c_e.alignment = Alignment(horizontal="right")

            c_g = ws.cell(curr_r, 7, l.libur)
            c_g.font = f_reg_9
            c_g.alignment = Alignment(horizontal="right")

            c_h = ws.cell(curr_r, 8, f"=D{curr_r}+E{curr_r}+F{curr_r}+G{curr_r}")
            c_h.font = f_reg_9
            c_h.number_format = num_fmt_int
            c_h.alignment = Alignment(horizontal="right")
            c_h.border = Border(right=med_side)

            if idx == len(lines) - 1:
                for ci in range(4, 9):
                    ws.cell(curr_r, ci).border = Border(bottom=double_side, right=med_side if ci == 8 else None)

        plot_sum_row = plot_start_row + num_pos
        ws.cell(plot_sum_row, 3).border = Border(left=med_side, bottom=med_side)
        for ci, col_let in [(4, 'D'), (5, 'E'), (6, 'F'), (7, 'G'), (8, 'H')]:
            c = ws.cell(plot_sum_row, ci, f"=SUM({col_let}{plot_start_row}:{col_let}{plot_sum_row-1})")
            c.font = f_bold_9
            c.alignment = Alignment(horizontal="right")
            if ci == 8:
                c.number_format = num_fmt_int
            c.border = Border(bottom=med_side, right=med_side if ci == 8 else None)

        # ----------------------------------------------------
        # 2. SECTION BIAYA SDM
        # ----------------------------------------------------
        biaya_title_row = plot_sum_row + 2
        ws.cell(biaya_title_row, 3, "BIAYA SDM").font = f_sec_hdr
        for ci in range(3, 9):
            c = ws.cell(biaya_title_row, ci)
            c.fill = fill_yellow
            c.border = Border(top=med_side, left=med_side if ci == 3 else None, right=med_side if ci == 8 else None)

        biaya_hdr_row = biaya_title_row + 1
        ws.cell(biaya_hdr_row, 6, "UNIT ").font = f_bold_11
        ws.cell(biaya_hdr_row, 6).alignment = Alignment(horizontal="right")
        ws.cell(biaya_hdr_row, 7, "UNIT COST").font = f_bold_11
        ws.cell(biaya_hdr_row, 7).alignment = Alignment(horizontal="right")
        ws.cell(biaya_hdr_row, 8, "TOTAL").font = f_bold_11
        ws.cell(biaya_hdr_row, 8).alignment = Alignment(horizontal="right")
        ws.cell(biaya_hdr_row, 8).number_format = num_fmt_curr
        ws.cell(biaya_hdr_row, 3).border = Border(left=med_side)
        ws.cell(biaya_hdr_row, 8).border = Border(right=med_side)

        biaya_start_row = biaya_hdr_row + 1
        biaya_sum_row = biaya_start_row + num_pos

        # Hitung Offset Baris untuk Section 3
        sec3_title_row = biaya_sum_row + 2
        sec3_hdr_row = sec3_title_row + 2
        sec3_gapok_row = sec3_hdr_row + 2
        sec3_tunjangan_row = sec3_gapok_row + 1
        sec3_upah_tot_row = sec3_tunjangan_row + 1
        sec3_ovh_title_row = sec3_upah_tot_row + 2
        sec3_bpjs_kes_row = sec3_ovh_title_row + 1
        sec3_bpjs_tk_row = sec3_bpjs_kes_row + 1
        sec3_seragam_row = sec3_bpjs_tk_row + 1
        sec3_thr_row = sec3_seragam_row + 1
        sec3_overhead_row = sec3_thr_row + 1
        sec3_fee_row = sec3_overhead_row + 1
        sec3_ovh_tot_row = sec3_fee_row + 1
        sec3_unit_cost_row = sec3_ovh_tot_row + 2

        note_title_row = sec3_unit_cost_row + 3
        note_val_row = note_title_row + 1
        umk_cell_ref = f"$D${note_val_row}"

        for idx, l in enumerate(lines):
            curr_r = biaya_start_row + idx
            plot_r = plot_start_row + idx
            pos_col_letter = get_column_letter(7 + idx)

            ws.cell(curr_r, 3, l.jabatan).font = f_reg_11
            ws.cell(curr_r, 3).border = Border(left=med_side)

            c_f = ws.cell(curr_r, 6, f"=H{plot_r}")
            c_f.font = f_reg_11
            c_f.alignment = Alignment(horizontal="right")

            c_g = ws.cell(curr_r, 7, f"={pos_col_letter}{sec3_unit_cost_row}")
            c_g.font = f_reg_11
            c_g.number_format = num_fmt_curr

            c_h = ws.cell(curr_r, 8, f"=G{curr_r}*F{curr_r}")
            c_h.font = f_reg_11
            c_h.number_format = num_fmt_curr
            c_h.border = Border(right=med_side)

            if idx == len(lines) - 1:
                for ci in [6, 7, 8]:
                    ws.cell(curr_r, ci).border = Border(bottom=double_side, right=med_side if ci == 8 else None)

        ws.cell(biaya_sum_row, 3).border = Border(left=med_side, bottom=med_side)
        for ci in range(4, 9):
            ws.cell(biaya_sum_row, ci).border = Border(bottom=med_side, right=med_side if ci == 8 else None)
        c_tot_unit = ws.cell(biaya_sum_row, 6, f"=SUM(F{biaya_start_row}:F{biaya_sum_row-1})")
        c_tot_unit.font = f_bold_11
        c_tot_unit.alignment = Alignment(horizontal="right")

        c_tot_biaya = ws.cell(biaya_sum_row, 8, f"=SUM(H{biaya_start_row}:H{biaya_sum_row-1})")
        c_tot_biaya.font = f_bold_11
        c_tot_biaya.number_format = num_fmt_curr

        # ----------------------------------------------------
        # 3. SECTION RINCIAN UNIT COST SDM
        # ----------------------------------------------------
        ws.cell(sec3_title_row, 3, "RINCIAN UNIT COST SDM").font = f_title
        ws.cell(sec3_hdr_row, 3, "JABATAN SDM").font = f_sec_hdr

        max_sec3_col = max(8, 6 + num_pos)
        for ci in range(3, max_sec3_col + 1):
            c = ws.cell(sec3_hdr_row, ci)
            c.fill = fill_yellow
            c.border = Border(top=med_side, left=med_side if ci == 3 else None, right=med_side if ci == max_sec3_col else None)

        for idx, l in enumerate(lines):
            pos_col = 7 + idx
            j_name = l.jabatan or ''
            code = 'TL' if 'LEADER' in j_name.upper() or 'TL' in j_name.upper() else ('CSO' if 'CSO' in j_name.upper() else j_name)
            c = ws.cell(sec3_hdr_row, pos_col, code)
            c.font = f_bold_11
            c.alignment = Alignment(horizontal="center")

        ws.cell(sec3_hdr_row + 1, 3, "KOMPONEN UPAH TETAP").font = f_reg_11
        ws.cell(sec3_hdr_row + 1, 3).border = Border(left=med_side)
        ws.cell(sec3_hdr_row + 1, max_sec3_col).border = Border(right=med_side)

        ws.cell(sec3_gapok_row, 3, "     GAPOK").font = f_reg_11
        ws.cell(sec3_gapok_row, 3).border = Border(left=med_side)
        ws.cell(sec3_gapok_row, max_sec3_col).border = Border(right=med_side)
        for idx, l in enumerate(lines):
            pos_col = 7 + idx
            c = ws.cell(sec3_gapok_row, pos_col, l.gapok)
            c.font = f_reg_11
            c.number_format = num_fmt_curr
            c.alignment = Alignment(horizontal="right")

        ws.cell(sec3_tunjangan_row, 3, "     TUNJANGAN JABATAN ").font = f_reg_11
        ws.cell(sec3_tunjangan_row, 3).border = Border(left=med_side)
        ws.cell(sec3_tunjangan_row, max_sec3_col).border = Border(right=med_side)
        for idx, l in enumerate(lines):
            pos_col = 7 + idx
            tunj = l.tunjangan
            c = ws.cell(sec3_tunjangan_row, pos_col, tunj if tunj > 0 else '-')
            c.font = f_bold_11 if tunj <= 0 or tunj == '-' else f_reg_11
            c.number_format = num_fmt_curr
            c.alignment = Alignment(horizontal="right")

        ws.merge_cells(f"D{sec3_upah_tot_row}:E{sec3_upah_tot_row}")
        ws.cell(sec3_upah_tot_row, 4, "TOTAL").font = f_bold_11
        ws.cell(sec3_upah_tot_row, 4).alignment = Alignment(horizontal="center")
        ws.cell(sec3_upah_tot_row, 3).border = Border(left=med_side, bottom=med_side)
        for ci in range(4, max_sec3_col + 1):
            ws.cell(sec3_upah_tot_row, ci).border = Border(bottom=med_side, right=med_side if ci == max_sec3_col else None)

        for idx, l in enumerate(lines):
            pos_col = 7 + idx
            col_let = get_column_letter(pos_col)
            c = ws.cell(sec3_upah_tot_row, pos_col, f"=SUM({col_let}{sec3_gapok_row}:{col_let}{sec3_tunjangan_row})")
            c.font = f_bold_11
            c.number_format = num_fmt_curr
            c.alignment = Alignment(horizontal="right")

        ws.cell(sec3_ovh_title_row, 3, "BIAYA OVERHEAD").font = f_sec_hdr
        for ci in range(3, max_sec3_col + 1):
            c = ws.cell(sec3_ovh_title_row, ci)
            c.fill = fill_yellow
            c.border = Border(top=med_side, left=med_side if ci == 3 else None, right=med_side if ci == max_sec3_col else None)

        # BPJS KESEHATAN
        ws.cell(sec3_bpjs_kes_row, 3, "BPJS KESEHATAN").font = f_reg_11
        ws.cell(sec3_bpjs_kes_row, 3).border = Border(left=med_side)
        ws.cell(sec3_bpjs_kes_row, max_sec3_col).border = Border(right=med_side)
        ws.cell(sec3_bpjs_kes_row, 5, self.bpjs_kes_rate or 0.04).font = f_bold_11
        ws.cell(sec3_bpjs_kes_row, 5).number_format = "0.0%"
        ws.cell(sec3_bpjs_kes_row, 5).alignment = Alignment(horizontal="center")
        for idx, l in enumerate(lines):
            pos_col = 7 + idx
            c = ws.cell(sec3_bpjs_kes_row, pos_col, f"=E{sec3_bpjs_kes_row}*{umk_cell_ref}")
            c.font = f_reg_11
            c.number_format = num_fmt_curr
            c.alignment = Alignment(horizontal="right")

        # BPJS KETENAGAKERJAAN
        ws.cell(sec3_bpjs_tk_row, 3, "BPJS KETENAGAKERJAAN").font = f_reg_11
        ws.cell(sec3_bpjs_tk_row, 3).border = Border(left=med_side)
        ws.cell(sec3_bpjs_tk_row, max_sec3_col).border = Border(right=med_side)
        ws.cell(sec3_bpjs_tk_row, 5, self.bpjs_tk_rate or 0.0689).font = f_bold_11
        ws.cell(sec3_bpjs_tk_row, 5).number_format = "0.00%"
        ws.cell(sec3_bpjs_tk_row, 5).alignment = Alignment(horizontal="center")
        for idx, l in enumerate(lines):
            pos_col = 7 + idx
            c = ws.cell(sec3_bpjs_tk_row, pos_col, f"={umk_cell_ref}*E{sec3_bpjs_tk_row}")
            c.font = f_reg_11
            c.number_format = num_fmt_curr
            c.alignment = Alignment(horizontal="right")

        # SERAGAM
        ws.cell(sec3_seragam_row, 3, "SERAGAM").font = f_reg_11
        ws.cell(sec3_seragam_row, 3).border = Border(left=med_side)
        ws.cell(sec3_seragam_row, max_sec3_col).border = Border(right=med_side)
        for idx, l in enumerate(lines):
            pos_col = 7 + idx
            c = ws.cell(sec3_seragam_row, pos_col, l.seragam)
            c.font = f_reg_11
            c.number_format = num_fmt_curr
            c.alignment = Alignment(horizontal="right")

        # THR
        ws.cell(sec3_thr_row, 3, "THR").font = f_reg_11
        ws.cell(sec3_thr_row, 3).border = Border(left=med_side)
        ws.cell(sec3_thr_row, max_sec3_col).border = Border(right=med_side)
        ws.cell(sec3_thr_row, 5, 12).font = f_bold_11
        ws.cell(sec3_thr_row, 5).alignment = Alignment(horizontal="center")
        for idx, l in enumerate(lines):
            pos_col = 7 + idx
            col_let = get_column_letter(pos_col)
            c = ws.cell(sec3_thr_row, pos_col, f"={col_let}{sec3_upah_tot_row}/E{sec3_thr_row}")
            c.font = f_reg_11
            c.number_format = num_fmt_curr
            c.alignment = Alignment(horizontal="right")

        # OVERHEAD (2%)
        ws.cell(sec3_overhead_row, 3, "Overhead").font = f_reg_11
        ws.cell(sec3_overhead_row, 3).border = Border(left=med_side)
        ws.cell(sec3_overhead_row, max_sec3_col).border = Border(right=med_side)
        ws.cell(sec3_overhead_row, 5, 0.02).font = f_bold_11
        ws.cell(sec3_overhead_row, 5).number_format = "0%"
        ws.cell(sec3_overhead_row, 5).alignment = Alignment(horizontal="center")
        for idx, l in enumerate(lines):
            pos_col = 7 + idx
            col_let = get_column_letter(pos_col)
            c = ws.cell(sec3_overhead_row, pos_col, f"=E{sec3_overhead_row}*{col_let}{sec3_upah_tot_row}")
            c.font = f_reg_11
            c.number_format = num_fmt_curr
            c.alignment = Alignment(horizontal="right")

        # FEE (2%)
        ws.cell(sec3_fee_row, 3, "FEE").font = f_reg_11
        ws.cell(sec3_fee_row, 3).border = Border(left=med_side)
        ws.cell(sec3_fee_row, max_sec3_col).border = Border(right=med_side)
        ws.cell(sec3_fee_row, 5, 0.02).font = f_bold_11
        ws.cell(sec3_fee_row, 5).number_format = "0%"
        ws.cell(sec3_fee_row, 5).alignment = Alignment(horizontal="center")
        for idx, l in enumerate(lines):
            pos_col = 7 + idx
            col_let = get_column_letter(pos_col)
            c = ws.cell(sec3_fee_row, pos_col, f"=E{sec3_fee_row}*{col_let}{sec3_upah_tot_row}")
            c.font = f_reg_11
            c.number_format = num_fmt_curr
            c.alignment = Alignment(horizontal="right")

        # TOTAL OVERHEAD
        ws.merge_cells(f"D{sec3_ovh_tot_row}:E{sec3_ovh_tot_row}")
        ws.cell(sec3_ovh_tot_row, 4, "TOTAL").font = f_bold_11
        ws.cell(sec3_ovh_tot_row, 4).alignment = Alignment(horizontal="center")
        ws.cell(sec3_ovh_tot_row, 3).border = Border(left=med_side, bottom=med_side)
        for ci in range(4, max_sec3_col + 1):
            ws.cell(sec3_ovh_tot_row, ci).border = Border(bottom=med_side, right=med_side if ci == max_sec3_col else None)

        for idx, l in enumerate(lines):
            pos_col = 7 + idx
            col_let = get_column_letter(pos_col)
            c = ws.cell(sec3_ovh_tot_row, pos_col, f"=SUM({col_let}{sec3_bpjs_kes_row}:{col_let}{sec3_fee_row})")
            c.font = f_bold_11
            c.number_format = num_fmt_curr
            c.alignment = Alignment(horizontal="right")

        # TOTAL UNIT COST
        ws.cell(sec3_unit_cost_row, 3, "TOTAL UNIT COST").font = f_bold_11
        for ci in range(3, max_sec3_col + 1):
            c = ws.cell(sec3_unit_cost_row, ci)
            c.fill = fill_yellow
            c.border = Border(top=med_side, bottom=med_side, left=med_side if ci == 3 else None, right=med_side if ci == max_sec3_col else None)

        for idx, l in enumerate(lines):
            pos_col = 7 + idx
            col_let = get_column_letter(pos_col)
            c = ws.cell(sec3_unit_cost_row, pos_col, f"={col_let}{sec3_upah_tot_row}+{col_let}{sec3_ovh_tot_row}")
            c.font = f_bold_11
            c.number_format = num_fmt_curr
            c.alignment = Alignment(horizontal="right")

        # ----------------------------------------------------
        # 4. SECTION NOTE & UMK ACUAN
        # ----------------------------------------------------
        ws.cell(note_title_row, 2, "Note").font = f_reg_11
        ws.cell(note_val_row, 3, self.keterangan_umk or "1. Nilai BPJS mengunakan UMK").font = f_reg_11
        c_umk = ws.cell(note_val_row, 4, self.dasar_umk or 2400000.0)
        c_umk.font = f_reg_11
        c_umk.number_format = num_fmt_curr

        # Simpan & Return Attachment
        fp = io.BytesIO()
        wb.save(fp)
        data_base64 = base64.b64encode(fp.getvalue())

        attachment = self.env['ir.attachment'].create({
            'name': f'DETAIL_SDM_{self.name.replace("/", "_")}.xlsx',
            'type': 'binary',
            'datas': data_base64,
            'res_model': self._name,
            'res_id': self.id,
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        })
        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }


class KsgHcManpowerCostSheetLine(models.Model):
    _name = 'ksg.hc.manpower.cost.sheet.line'
    _description = 'Pos Rincian SDM'

    cost_sheet_id = fields.Many2one('ksg.hc.manpower.cost.sheet', string='Dokumen SDM Cost', ondelete='cascade')
    currency_id = fields.Many2one(related='cost_sheet_id.currency_id')

    jabatan = fields.Char(string='Jabatan SDM', required=True)
    shift_1 = fields.Integer(string='Shift 1', default=0)
    shift_2 = fields.Integer(string='Shift 2', default=0)
    libur = fields.Integer(string='Libur', default=0)
    unit_tk = fields.Integer(string='Unit TK (Org)', default=1)

    gapok = fields.Monetary(string='Gaji Pokok', currency_field='currency_id')
    tunjangan = fields.Monetary(string='Tunjangan Jabatan', currency_field='currency_id')
    total_upah_tetap = fields.Monetary(string='Total Upah Tetap', compute='_compute_cost', store=True, currency_field='currency_id')

    bpjs_kes = fields.Monetary(string='BPJS Kesehatan', currency_field='currency_id')
    bpjs_tk = fields.Monetary(string='BPJS TK', currency_field='currency_id')
    seragam = fields.Monetary(string='Seragam', currency_field='currency_id')
    thr = fields.Monetary(string='THR', compute='_compute_cost', store=True, currency_field='currency_id')
    overhead = fields.Monetary(string='Overhead (2%)', compute='_compute_cost', store=True, currency_field='currency_id')
    fee = fields.Monetary(string='Fee (2%)', compute='_compute_cost', store=True, currency_field='currency_id')
    total_overhead = fields.Monetary(string='Total Overhead', compute='_compute_cost', store=True, currency_field='currency_id')

    unit_cost = fields.Monetary(string='Total Unit Cost', compute='_compute_cost', store=True, currency_field='currency_id')
    total_biaya = fields.Monetary(string='Total Tagihan/Bln', compute='_compute_cost', store=True, currency_field='currency_id')

    @api.depends('unit_tk', 'gapok', 'tunjangan', 'bpjs_kes', 'bpjs_tk', 'seragam')
    def _compute_cost(self):
        for line in self:
            upah = line.gapok + line.tunjangan
            line.total_upah_tetap = upah
            line.thr = upah / 12.0
            line.overhead = 0.02 * upah
            line.fee = 0.02 * upah
            ovh = line.bpjs_kes + line.bpjs_tk + line.seragam + line.thr + line.overhead + line.fee
            line.total_overhead = ovh
            u_cost = upah + ovh
            line.unit_cost = u_cost
            line.total_biaya = u_cost * line.unit_tk


class KsgHcAssignmentRequest(models.Model):
    _name = 'ksg.hc.assignment.request'
    _description = 'Penetapan Assignment Pekerja oleh HC'
    _inherit = ['mail.thread', 'mail.activity.mixin']

    name = fields.Char(string='Nomor Penetapan', required=True, copy=False, default=lambda self: _('New'))
    project_id = fields.Many2one('ksg.sales.project', string='Proyek Terkait', required=True, tracking=True)
    worker_id = fields.Many2one('ksg.sales.worker', string='Tenaga Kerja', required=True, tracking=True)
    sumber_kebutuhan = fields.Selection([
        ('penjualan', 'Penjualan'),
        ('operasional', 'Operasional')
    ], string='Sumber Kebutuhan', default='operasional', required=True)
    ditetapkan_oleh = fields.Selection([('hc', 'Human Capital (HC)')], default='hc', readonly=True)
    tanggal_assign = fields.Date(string='Tanggal Penugasan', default=fields.Date.context_today)
    tanggal_selesai = fields.Date(string='Tanggal Berakhir')
    state = fields.Selection([('draft', 'Draft'), ('active', 'Aktif'), ('ended', 'Selesai')], default='draft')

    def action_assign(self):
        self.ensure_one()
        self.write({'state': 'active'})
        self.env['ksg.operational.assignment'].create({
            'project_id': self.project_id.id,
            'worker_name': getattr(self.worker_id, 'name', 'Pekerja'),
            'jabatan': getattr(self.worker_id, 'posisi', '') or 'Pekerja',
            'tanggal_mulai': self.tanggal_assign,
            'tanggal_selesai': self.tanggal_selesai,
            'status': 'active'
        })


class KsgHcRecruitment(models.Model):
    _name = 'ksg.hc.recruitment'
    _description = 'Proses Rekrutmen & e-Sign PKWT'
    _inherit = ['mail.thread', 'mail.activity.mixin']

    name = fields.Char(string='ID Rekrutmen', required=True, copy=False, default=lambda self: _('New'))
    project_id = fields.Many2one('ksg.sales.project', string='Proyek Terkait')
    posisi = fields.Char(string='Posisi Lowongan', required=True)
    kandidat_name = fields.Char(string='Nama Kandidat / Pelamar', required=True)
    referensi_ksg_career = fields.Char(string='ID / Tautan KSG Career')
    status_seleksi = fields.Selection([
        ('screening', 'Screening Berkas'),
        ('interview', 'Interview / Wawancara'),
        ('offering', 'Offering Letter'),
        ('accepted', 'Diterima & TTD PKWT'),
        ('rejected', 'Ditolak')
    ], string='Status Seleksi', default='screening', tracking=True)
    dokumen_kontrak = fields.Binary(string='Berkas Kontrak PKWT')
    filename_kontrak = fields.Char(string='Nama File Kontrak')
    tanggal_ttd_digital = fields.Date(string='Tanggal TTD Digital')


class KsgHcPayrollSync(models.Model):
    _name = 'ksg.hc.payroll.sync'
    _description = 'Integrasi & Approval Rekap Payroll HC'
    _inherit = ['mail.thread', 'mail.activity.mixin']

    name = fields.Char(string='Nomor Rekap Payroll', required=True, default=lambda self: _('New'))
    periode = fields.Char(string='Periode (Bulan-Tahun)', required=True, default='09-2026')
    status_kirim = fields.Selection([('belum_kirim', 'Belum Kirim'), ('terkirim', 'Terkirim ke Keuangan'), ('gagal', 'Gagal')], default='belum_kirim')
    file_excel_backup = fields.Binary(string='File Excel Rekap Payroll (Wajib Cross-Check)', required=True)
    filename_excel = fields.Char(string='Nama File Excel')
    total_gaji_diajukan = fields.Monetary(string='Total Gaji Diajukan (Rp)', currency_field='currency_id')
    currency_id = fields.Many2one('res.currency', default=lambda self: self.env.company.currency_id)
    state = fields.Selection([
        ('draft', 'Draft HC'),
        ('diajukan_hc', 'Menunggu Approval Direktur'),
        ('disetujui_direktur', 'Disetujui Direktur'),
        ('dikirim_keuangan', 'Diteruskan ke Keuangan')
    ], default='draft', tracking=True)

    def action_submit_hc(self):
        self.write({'state': 'diajukan_hc'})

    def action_approve_direktur(self):
        self.write({'state': 'disetujui_direktur', 'status_kirim': 'terkirim'})


class KsgHcExitClearance(models.Model):
    _name = 'ksg.hc.exit.clearance'
    _description = 'Proses Exit Clearance Karyawan'
    _inherit = ['mail.thread', 'mail.activity.mixin']

    name = fields.Char(string='Nomor Exit Clearance', default=lambda self: _('New'))
    worker_id = fields.Many2one('ksg.sales.worker', string='Karyawan / Pekerja', required=True)
    tanggal_pengajuan_resign = fields.Date(string='Tanggal Pengajuan Resign', default=fields.Date.context_today, required=True)
    tanggal_efektif_resign = fields.Date(string='Tanggal Efektif Resign', required=True)
    notice_period_valid = fields.Boolean(string='Notice Period Valid (H-30)', compute='_compute_notice_period', store=True)
    status = fields.Selection([('draft', 'Draft'), ('berjalan', 'Sedang Berjalan'), ('selesai', 'Selesai')], default='draft', tracking=True)
    checklist_line_ids = fields.One2many('ksg.hc.exit.clearance.line', 'clearance_id', string='Checklist Serah Terima')

    @api.depends('tanggal_pengajuan_resign', 'tanggal_efektif_resign')
    def _compute_notice_period(self):
        for rec in self:
            if rec.tanggal_pengajuan_resign and rec.tanggal_efektif_resign:
                delta = (rec.tanggal_efektif_resign - rec.tanggal_pengajuan_resign).days
                rec.notice_period_valid = (delta >= 30)
            else:
                rec.notice_period_valid = False

    def action_start(self):
        for rec in self:
            if not rec.notice_period_valid:
                raise ValidationError(_("Pengajuan resign wajib minimal 30 hari (H-30) sebelum tanggal efektif."))
            rec.status = 'berjalan'
            if rec.worker_id:
                rec.worker_id.status_resign = True
                rec.worker_id.kompensasi_locked = True

    def action_done(self):
        for rec in self:
            rec.status = 'selesai'
            if rec.worker_id:
                rec.worker_id.kompensasi_locked = False


class KsgHcExitClearanceLine(models.Model):
    _name = 'ksg.hc.exit.clearance.line'
    _description = 'Baris Checklist Exit Clearance'

    clearance_id = fields.Many2one('ksg.hc.exit.clearance', ondelete='cascade')
    item = fields.Char(string='Item Serah Terima (Pekerjaan / Aset / ID Card)', required=True)
    status = fields.Selection([('pending', 'Pending'), ('selesai', 'Selesai Serah Terima')], default='pending')
    penerima = fields.Char(string='Penerima Handover')
    tanggal_serah_terima = fields.Date(string='Tanggal Serah Terima')


class KsgHcSppdSync(models.Model):
    _name = 'ksg.hc.sppd.sync'
    _description = 'Data Sinkronisasi SPPD (Presenly)'

    name = fields.Char(string='Nomor SPPD', required=True)
    worker_id = fields.Many2one('ksg.sales.worker', string='Karyawan')
    source_system = fields.Char(string='Sumber', default='Presenly')
    tanggal_berangkat = fields.Date(string='Tanggal Berangkat')
    tanggal_kembali = fields.Date(string='Tanggal Kembali')
    tujuan = fields.Char(string='Kota / Lokasi Tujuan')
    keperluan = fields.Text(string='Maksud / Keperluan Dinas')
    status = fields.Char(string='Status SPPD', default='Approved')